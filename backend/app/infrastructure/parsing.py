"""Document text extraction (Phase 5, P5-3/P5-6).

Security and robustness posture:
- File type is decided by MAGIC BYTES, never by extension or the
  browser-supplied content type (both trivially forged).
- Every library exception is converted to a domain error at this boundary:
  corrupt, encrypted, or truncated files become DocumentParseError — a 500
  must never be reachable through upload content.
- Extraction is synchronous (P5-6): at the 5 MB cap, typical resumes parse
  in well under a second; background processing is a Phase 10 concern.
"""

import io
from typing import Literal

import pdfplumber
from docx import Document as DocxDocument

from app.domain.exceptions import (
    DocumentParseError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)

PDF_MAGIC = b"%PDF-"
DOCX_MAGIC = b"PK\x03\x04"  # ZIP local-file header — a DOCX is a ZIP container


def sniff_file_type(data: bytes) -> Literal["pdf", "docx"]:
    """Identify the real file type from its first bytes."""
    if data.startswith(PDF_MAGIC):
        return "pdf"
    if data.startswith(DOCX_MAGIC):
        return "docx"
    raise UnsupportedFileTypeError


def extract_text_pdf(data: bytes) -> str:
    """Extract text from all pages; raise domain errors on failure/emptiness."""
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
    except Exception as exc:
        # pdfminer raises a zoo of exception types (encryption, corruption,
        # truncation); one boundary conversion keeps them all out of the API.
        raise DocumentParseError from exc
    text = "\n".join(pages).strip()
    if not text:
        raise EmptyDocumentError
    return text


def extract_text_docx(data: bytes) -> str:
    """Extract text from paragraphs AND tables (resume templates use both)."""
    try:
        document = DocxDocument(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    except Exception as exc:
        raise DocumentParseError from exc
    text = "\n".join(parts).strip()
    if not text:
        raise EmptyDocumentError
    return text


def extract_text(data: bytes) -> str:
    """Dispatch on magic bytes. The only function the service layer calls."""
    file_type = sniff_file_type(data)
    if file_type == "pdf":
        return extract_text_pdf(data)
    return extract_text_docx(data)
