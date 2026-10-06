"""Resume upload, parsing, and ownership tests (Phase 5 Step 2).

Test-fixing strategy (no new dependencies, no binary fixtures):
- Real DOCX bytes are generated IN-TEST with python-docx (already pinned).
- A minimal structurally-valid PDF is assembled programmatically with a
  computed xref table — deterministic, ~25 lines.
- Corrupt/encrypted/garbage paths use hand-crafted byte sequences; the
  parsing boundary converts every library failure to one domain error.
"""

import io

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from httpx import Response

from app.domain.exceptions import (
    DocumentParseError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)
from app.infrastructure.parsing import extract_text, sniff_file_type

RESUMES = "/api/v1/resumes"

PASSWORD = "correct-horse-battery"

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
def _error_code(response: Response) -> str:
    body = response.json()
    assert set(body) == {"error"}, f"expected error envelope, got: {body}"
    return body["error"]["code"]


def _register_verify_login(
    client: TestClient, outbox, *, email: str = "user@example.com"
) -> None:
    client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": PASSWORD, "name": "Test User"},
    )
    body = outbox.messages[-1].body
    start = body.index("token=") + len("token=")
    token = body[start:].split()[0]
    client.post("/api/v1/auth/verify-email", json={"token": token})
    client.post(
        "/api/v1/auth/login", json={"email": email, "password": PASSWORD}
    )


def _csrf(client: TestClient) -> dict[str, str]:
    return {"X-CSRF-Token": client.cookies.get("csrf_token")}


def _build_minimal_pdf(text: str) -> bytes:
    """Assemble a tiny single-page PDF with a computed xref — structurally
    valid, so pdfplumber parses it without warnings."""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        (
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
            + stream + b"\nendstream"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def _build_docx(text: str) -> bytes:
    buffer = io.BytesIO()
    document = DocxDocument()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, *, name: str, content: bytes, mime: str) -> Response:
    return client.post(
        RESUMES,
        files={"file": (name, content, mime)},
        headers=_csrf(client),
    )


# --- parsing unit tests -------------------------------------------------------


def test_sniff_identifies_pdf_and_docx_by_magic_bytes():
    assert sniff_file_type(b"%PDF-1.7 whatever") == "pdf"
    assert sniff_file_type(b"PK\x03\x04 whatever") == "docx"


def test_sniff_rejects_unknown_magic():
    with pytest.raises(UnsupportedFileTypeError):
        sniff_file_type(b"GIF89a not a resume")
    with pytest.raises(UnsupportedFileTypeError):
        sniff_file_type(b"")


def test_extract_text_pdf_happy():
    text = extract_text(_build_minimal_pdf("Hello Resume"))
    assert "Hello Resume" in text


def test_extract_text_docx_happy():
    text = extract_text(_build_docx("Curriculum Vitae body text"))
    assert "Curriculum Vitae body text" in text


def test_extract_text_docx_empty_raises_empty_document():
    buffer = io.BytesIO()
    DocxDocument().save(buffer)  # structurally valid, zero paragraphs
    with pytest.raises(EmptyDocumentError):
        extract_text(buffer.getvalue())


def test_corrupt_pdf_raises_parse_error_not_library_exception():
    with pytest.raises(DocumentParseError):
        extract_text(b"%PDF-1.4 this is not really a pdf at all")


# --- API: upload ---------------------------------------------------------------


def test_upload_docx_returns_201_and_persists(client, email_outbox, db):
    _register_verify_login(client, email_outbox)

    response = _upload(
        client,
        name="my-resume.docx",
        content=_build_docx("Worked at Example Corp"),
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert response.status_code == 201
    body = response.json()
    assert body["filename"] == "my-resume.docx"
    assert body["status"] == "parsed"
    assert body["file_size"] > 0
    # Durable: the detail endpoint serves it back with extracted text.
    detail = client.get(f"{RESUMES}/{body['id']}")
    assert detail.status_code == 200
    assert "Worked at Example Corp" in detail.json()["raw_text"]


def test_upload_pdf_returns_201(client, email_outbox):
    _register_verify_login(client, email_outbox)

    response = _upload(
        client,
        name="resume.pdf",
        content=_build_minimal_pdf("PDF resume text"),
        mime="application/pdf",
    )

    assert response.status_code == 201
    assert response.json()["content_type"] == "application/pdf"


def test_upload_requires_authentication(client):
    response = client.post(
        RESUMES, files={"file": ("r.pdf", b"%PDF-1.4 x", "application/pdf")}
    )
    assert response.status_code == 401
    assert _error_code(response) == "not_authenticated"


def test_upload_requires_csrf(client, email_outbox):
    _register_verify_login(client, email_outbox)

    response = client.post(
        RESUMES, files={"file": ("r.pdf", b"%PDF-1.4 x", "application/pdf")}
    )

    assert response.status_code == 403
    assert _error_code(response) == "csrf_failed"


def test_upload_rejects_wrong_file_type(client, email_outbox):
    _register_verify_login(client, email_outbox)

    response = _upload(
        client, name="virus.exe", content=b"MZ fake executable", mime="application/octet-stream"
    )

    assert response.status_code == 415
    assert _error_code(response) == "unsupported_file_type"


def test_upload_rejects_oversized_files(client, email_outbox):
    from app.core.config import settings

    _register_verify_login(client, email_outbox)
    oversized = b"%PDF-1.4" + b"x" * settings.MAX_RESUME_SIZE_BYTES

    response = _upload(client, name="big.pdf", content=oversized, mime="application/pdf")

    assert response.status_code == 413
    assert _error_code(response) == "file_too_large"


def test_upload_rejects_corrupt_pdf_with_parse_error(client, email_outbox):
    _register_verify_login(client, email_outbox)

    response = _upload(
        client, name="broken.pdf", content=b"%PDF-1.4 garbage not a pdf", mime="application/pdf"
    )

    assert response.status_code == 422
    assert _error_code(response) == "document_parse_failed"


def test_upload_rejects_textless_docx_as_empty(client, email_outbox):
    _register_verify_login(client, email_outbox)
    buffer = io.BytesIO()
    DocxDocument().save(buffer)

    response = _upload(
        client, name="scanned.docx", content=buffer.getvalue(), mime=DOCX_MIME
    )

    assert response.status_code == 422
    assert _error_code(response) == "empty_document"


def test_upload_rejects_missing_file_field(client, email_outbox):
    _register_verify_login(client, email_outbox)

    response = client.post(RESUMES, headers=_csrf(client))

    assert response.status_code == 422
    assert _error_code(response) == "validation_error"


# --- API: list / get / delete ---------------------------------------------------


def test_list_returns_only_the_callers_resumes(client, email_outbox):
    _register_verify_login(client, email_outbox)
    _upload(client, name="a.docx", content=_build_docx("A resume"), mime=DOCX_MIME)
    _upload(client, name="b.docx", content=_build_docx("B resume"), mime=DOCX_MIME)

    response = client.get(RESUMES)

    assert response.status_code == 200
    items = response.json()
    assert len(items) == 2
    assert {item["filename"] for item in items} == {"a.docx", "b.docx"}
    # raw_text is deliberately absent from list items (large field):
    assert all("raw_text" not in item for item in items)


def test_get_detail_includes_extracted_text(client, email_outbox):
    _register_verify_login(client, email_outbox)
    created = _upload(
        client, name="r.docx", content=_build_docx("Detail text probe"), mime=DOCX_MIME
    ).json()

    detail = client.get(f"{RESUMES}/{created['id']}")

    assert detail.status_code == 200
    assert "Detail text probe" in detail.json()["raw_text"]


def test_foreign_resume_id_returns_404_idor_read(client, email_outbox):
    """THE IDOR TEST (owed since Step 8/S8-1): user B requests user A's
    resume id. 404 — never 403 (which would confirm existence), never 200."""
    _register_verify_login(client, email_outbox)  # user A
    created = _upload(
        client,
        name="secret.docx",
        content=_build_docx("A secret resume"),
        mime=DOCX_MIME,
    ).json()
    session_a = client.cookies.get("resumeai_session")

    _register_verify_login(
        client, email_outbox, email="attacker@example.com"
    )  # user B owns the jar  # user B owns the jar

    response = client.get(f"{RESUMES}/{created['id']}")

    assert response.status_code == 404
    assert _error_code(response) == "resume_not_found"
    # And the cookie juggling proves it: restore A, the resume is still there.
    client.cookies.set("resumeai_session", session_a)
    assert client.get(f"{RESUMES}/{created['id']}").status_code == 200


def test_foreign_delete_returns_404_and_target_survives_idor_delete(client, email_outbox):
    _register_verify_login(client, email_outbox)  # user A
    created = _upload(
        client,
        name="keepme.docx",
        content=_build_docx("Must survive"),
        mime=DOCX_MIME,
    ).json()
    session_a = client.cookies.get("resumeai_session")

    _register_verify_login(client, email_outbox, email="deleter@example.com")  # user B

    response = client.delete(f"{RESUMES}/{created['id']}", headers=_csrf(client))

    assert response.status_code == 404
    assert _error_code(response) == "resume_not_found"
    # The target SURVIVED the foreign delete attempt:
    client.cookies.set("resumeai_session", session_a)
    assert client.get(f"{RESUMES}/{created['id']}").status_code == 200


def test_delete_then_get_404s(client, email_outbox):
    _register_verify_login(client, email_outbox)
    created = _upload(
        client,
        name="todelete.docx",
        content=_build_docx("Delete me"),
        mime=DOCX_MIME,
    ).json()

    response = client.delete(f"{RESUMES}/{created['id']}", headers=_csrf(client))

    assert response.status_code == 204
    assert client.get(f"{RESUMES}/{created['id']}").status_code == 404


def test_invalid_uuid_path_returns_validation_error(client, email_outbox):
    _register_verify_login(client, email_outbox)

    response = client.get(f"{RESUMES}/not-a-uuid")

    assert response.status_code == 422
    assert _error_code(response) == "validation_error"
